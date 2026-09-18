from pathlib import Path

import core
import report


def test_compute_hash_rows_mixes_found_and_missing_files(tmp_path):
    (tmp_path / "a.pdf").write_bytes(b"conteudo")
    entries = [
        core.FileEntry(file_name="a.pdf", file_link="https://example.com/a", sha256_expected=None),
        core.FileEntry(file_name="faltando.zip", file_link="https://example.com/b", sha256_expected=None),
    ]

    rows = report.compute_hash_rows(entries, tmp_path)

    assert rows[0].file_name == "a.pdf"
    assert rows[0].sha256 == core.sha256_of_file(tmp_path / "a.pdf")
    assert rows[0].size_bytes == len(b"conteudo")

    assert rows[1].file_name == "faltando.zip"
    assert rows[1].sha256 is None
    assert rows[1].size_bytes == 0


def test_compute_hash_rows_cached_reuses_valid_cache_and_computes_the_rest(tmp_path):
    cached_path = tmp_path / "cached.pdf"
    cached_path.write_bytes(b"ja verificado no pipeline")
    stat = cached_path.stat()
    pending_path = tmp_path / "pending.pdf"
    pending_path.write_bytes(b"nunca calculado ainda")

    cached_entry = core.FileEntry(
        file_name="cached.pdf", file_link="https://example.com/a", sha256_expected=None,
        sha256_computed="cafecafe", sha256_computed_stat=(stat.st_size, stat.st_mtime),
    )
    pending_entry = core.FileEntry(file_name="pending.pdf", file_link="https://example.com/b", sha256_expected=None)

    progress_calls: list[str] = []
    rows = report.compute_hash_rows_cached(
        [cached_entry, pending_entry], tmp_path,
        on_progress=lambda entry, done, total, phase: progress_calls.append(entry.file_name),
    )

    by_name = {row.file_name: row for row in rows}
    assert by_name["cached.pdf"].sha256 == "cafecafe"
    assert by_name["pending.pdf"].sha256 == core.sha256_of_file(pending_path)
    # so o arquivo pendente passa pelo calculo (e reporta progresso) — o cacheado eh reaproveitado.
    assert "cached.pdf" not in progress_calls


def test_compute_hash_rows_cached_recomputes_when_file_changed_after_cache(tmp_path):
    path = tmp_path / "a.pdf"
    path.write_bytes(b"conteudo original")
    stat = path.stat()
    entry = core.FileEntry(
        file_name="a.pdf", file_link="https://example.com/a", sha256_expected=None,
        sha256_computed="hashantigo", sha256_computed_stat=(stat.st_size, stat.st_mtime),
    )
    path.write_bytes(b"conteudo mudou depois do cache")

    rows = report.compute_hash_rows_cached([entry], tmp_path)

    assert rows[0].sha256 == core.sha256_of_file(path)
    assert rows[0].sha256 != "hashantigo"


def test_humanize_bytes():
    assert report.humanize_bytes(0) == "0 bytes"
    assert report.humanize_bytes(1023) == "1023 bytes"
    assert report.humanize_bytes(1024) == "1.00 KB"
    assert report.humanize_bytes(1024 * 1024) == "1.00 MB"
    assert report.humanize_bytes(1024 * 1024 * 1024) == "1.00 GB"


def test_render_report_fills_placeholders_when_data_missing():
    text = report.render_report(report.ReportData())

    assert "[NÚMERO DO PROCESSO JUDICIAL]" in text
    assert "[VOLUME]" in text
    assert "[QUANTIDADE]" in text
    assert "(  ) Portal do Provedor" in text
    assert "| [nome-do-arquivo] | [hash] |" in text
    assert "**Provedor de Aplicação:** Apple Inc." in text
    assert "**Algoritmo utilizado:** SHA-256" in text


def test_render_report_fills_real_values_and_marks_checked_boxes():
    rows = [
        report.HashRow(file_name="a.pdf", sha256="deadbeef", size_bytes=1024),
        report.HashRow(file_name="faltando.zip", sha256=None, size_bytes=0),
    ]
    data = report.ReportData(
        numero_processo="0001234-56.2026.8.26.0000",
        recebido_portal=True,
        disponibiliza_processada=True,
        hash_rows=rows,
    )

    text = report.render_report(data)

    assert "0001234-56.2026.8.26.0000" in text
    assert "**Quantidade de Arquivos:** 2" in text
    assert "**Volume Total (Bytes/GB/TB):** 1.00 KB" in text
    assert "| a.pdf | deadbeef |" in text
    assert "| faltando.zip | (arquivo nao encontrado no destino) |" in text
    assert "(X) Portal do Provedor" in text
    assert "(  ) E-mail Oficial" in text
    assert "(X) Cópia Processada/Indexada" in text
    assert "(  ) Cópia da Aquisição Forense Original" in text


def test_render_report_includes_gpg_password_only_when_flagged():
    with_flag = report.render_report(
        report.ReportData(gpg_passphrase="Tr0ub4dor&3", incluir_senha_gpg=True)
    )
    assert "**Senha de Descriptografia (GPG):** Tr0ub4dor&3" in with_flag

    without_flag = report.render_report(
        report.ReportData(gpg_passphrase="Tr0ub4dor&3", incluir_senha_gpg=False)
    )
    assert "Senha de Descriptografia" not in without_flag

    flagged_but_empty = report.render_report(
        report.ReportData(gpg_passphrase="", incluir_senha_gpg=True)
    )
    assert "Senha de Descriptografia" not in flagged_but_empty


def test_render_report_html_includes_gpg_password_only_when_flagged():
    with_flag = report.render_report_html(
        report.ReportData(gpg_passphrase="Tr0ub4dor&3", incluir_senha_gpg=True)
    )
    assert "<strong>Senha de Descriptografia (GPG):</strong> Tr0ub4dor&amp;3" in with_flag

    without_flag = report.render_report_html(
        report.ReportData(gpg_passphrase="Tr0ub4dor&3", incluir_senha_gpg=False)
    )
    assert "Senha de Descriptografia" not in without_flag


def test_render_report_html_escapes_gpg_password():
    out = report.render_report_html(
        report.ReportData(gpg_passphrase="<script>1</script>", incluir_senha_gpg=True)
    )
    assert "<script>1</script>" not in out
    assert "&lt;script&gt;1&lt;/script&gt;" in out


def test_render_report_html_uses_arial_12_and_line_spacing_1_5():
    rows = [report.HashRow(file_name="a.pdf", sha256="deadbeef", size_bytes=1024)]
    data = report.ReportData(numero_processo="0001234", recebido_email=True, hash_rows=rows)

    out = report.render_report_html(data)

    assert "font-family:Arial" in out
    assert "font-size:12pt" in out
    assert "line-height:1.5" in out
    assert "<table" in out and "<th" in out
    assert "0001234" in out
    assert "☑ E-mail Oficial" in out  # checked box
    assert "☐ Portal do Provedor" in out  # unchecked box


def test_render_report_html_escapes_user_input():
    data = report.ReportData(numero_processo="<script>alert(1)</script> & cia")

    out = report.render_report_html(data)

    assert "<script>" not in out
    assert "&lt;script&gt;" in out
    assert "&amp; cia" in out


def test_build_docx_generates_a_valid_document_with_arial_12_and_spacing_1_5(tmp_path):
    from docx import Document
    from docx.shared import Pt

    rows = [
        report.HashRow(file_name="a.pdf", sha256="deadbeef", size_bytes=1024),
        report.HashRow(file_name="faltando.zip", sha256=None, size_bytes=0),
    ]
    data = report.ReportData(
        numero_processo="0001234-56.2026.8.26.0000",
        responsavel="Fulano de Tal",
        recebido_portal=True,
        hash_rows=rows,
    )

    out_path = tmp_path / "termo.docx"
    report.build_docx(data, out_path)

    assert out_path.is_file()

    doc = Document(str(out_path))
    normal = doc.styles["Normal"]
    assert normal.font.name == "Arial"
    assert normal.font.size == Pt(12)
    assert normal.paragraph_format.line_spacing == 1.5

    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "0001234-56.2026.8.26.0000" in full_text
    assert "Fulano de Tal" in full_text

    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert [c.text for c in table.rows[0].cells] == ["Arquivo", "Hash"]
    assert [c.text for c in table.rows[1].cells] == ["a.pdf", "deadbeef"]
    assert [c.text for c in table.rows[2].cells] == ["faltando.zip", "(arquivo nao encontrado no destino)"]
    assert "Senha de Descriptografia" not in full_text

    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            assert run.font.name == "Arial"


def test_build_docx_includes_gpg_password_only_when_flagged(tmp_path):
    from docx import Document

    data = report.ReportData(gpg_passphrase="Tr0ub4dor&3", incluir_senha_gpg=True)
    out_path = tmp_path / "termo_com_senha.docx"
    report.build_docx(data, out_path)

    full_text = "\n".join(p.text for p in Document(str(out_path)).paragraphs)
    assert "Senha de Descriptografia (GPG): Tr0ub4dor&3" in full_text

    data_sem_flag = report.ReportData(gpg_passphrase="Tr0ub4dor&3", incluir_senha_gpg=False)
    out_path_sem_flag = tmp_path / "termo_sem_senha.docx"
    report.build_docx(data_sem_flag, out_path_sem_flag)

    full_text_sem_flag = "\n".join(p.text for p in Document(str(out_path_sem_flag)).paragraphs)
    assert "Senha de Descriptografia" not in full_text_sem_flag
