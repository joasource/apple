"""
JoaKApple - gerador do Termo de Recebimento e Identificacao de Evidencia
Telematica.

Preenche o modelo fornecido pelo usuario com os dados digitados na GUI e com
os hashes SHA-256 calculados a partir dos arquivos efetivamente recebidos
(baixados) no diretorio de destino do pipeline.
"""

from __future__ import annotations

from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable, Optional

import core

FONT_NAME = "Arial"
FONT_SIZE_PT = 12
LINE_SPACING = 1.5


@dataclass
class HashRow:
    file_name: str
    sha256: Optional[str]  # None se o arquivo nao foi encontrado em disco
    size_bytes: int = 0


def compute_hash_rows(entries: list[core.FileEntry], output_dir: Path) -> list[HashRow]:
    """Calcula o SHA-256 de cada arquivo recebido (ja baixado) em output_dir."""
    rows: list[HashRow] = []
    for entry in entries:
        path = output_dir / entry.file_name
        if path.is_file():
            rows.append(HashRow(entry.file_name, core.sha256_of_file(path), path.stat().st_size))
        else:
            rows.append(HashRow(entry.file_name, None, 0))
    return rows


def compute_hash_rows_cached(
    entries: list[core.FileEntry],
    output_dir: Path,
    max_workers: int = 4,
    on_progress: Optional[Callable[[core.FileEntry, int, int, str], None]] = None,
) -> list[HashRow]:
    """Como compute_hash_rows, mas reaproveita o hash ja calculado na etapa de
    verificacao do pipeline (core.FileEntry.sha256_computed, valido enquanto o
    arquivo no disco nao mudar) e paraleliza o calculo do que sobrar, com
    progresso por arquivo via on_progress — evita recalcular do zero (serial e
    sem feedback) o que a GUI acabou de conferir."""
    rows: dict[str, HashRow] = {}
    pending: list[core.FileEntry] = []
    for entry in entries:
        path = output_dir / entry.file_name
        if not path.is_file():
            rows[entry.file_name] = HashRow(entry.file_name, None, 0)
            continue
        cached = core.cached_file_hash(entry, path)
        if cached is not None:
            rows[entry.file_name] = HashRow(entry.file_name, cached, path.stat().st_size)
        else:
            pending.append(entry)

    def _hash_one(entry: core.FileEntry) -> None:
        path = output_dir / entry.file_name
        progress = (
            (lambda done, total: on_progress(entry, done, total, "hash")) if on_progress else None
        )
        sha256 = core.sha256_of_file(path, on_progress=progress)
        rows[entry.file_name] = HashRow(entry.file_name, sha256, path.stat().st_size)

    if pending:
        with ThreadPoolExecutor(max_workers=max(1, max_workers), thread_name_prefix="report-hash") as pool:
            list(pool.map(_hash_one, pending))

    return [rows[e.file_name] for e in entries]


def humanize_bytes(n: int) -> str:
    size = float(n)
    for unit in ("bytes", "KB", "MB", "GB"):
        if size < 1024 or unit == "GB":
            return f"{int(size)} bytes" if unit == "bytes" else f"{size:.2f} {unit}"
        size /= 1024
    return f"{size:.2f} TB"


def _checkbox(checked: bool) -> str:
    return "(X)" if checked else "(  )"


def _checkbox_html(checked: bool) -> str:
    return "☑" if checked else "☐"  # ☑ / ☐


@dataclass
class ReportData:
    numero_processo: str = ""
    numero_pic: str = ""
    orgao_execucao: str = ""
    data: str = ""
    hora: str = ""
    responsavel: str = ""
    matricula: str = ""
    recebido_portal: bool = False
    recebido_email: bool = False
    procedimento_referenciado: str = ""
    id_documento: str = ""
    disponibiliza_original: bool = False
    disponibiliza_processada: bool = False
    hash_rows: list[HashRow] = field(default_factory=list)


def _summary(data: ReportData) -> tuple[str, str, int]:
    """Retorna (quantidade, volume, total_files), usando placeholders se nao houver arquivos."""
    total_files = len(data.hash_rows)
    total_bytes = sum(row.size_bytes for row in data.hash_rows)
    volume = humanize_bytes(total_bytes) if total_files else "[VOLUME]"
    quantidade = str(total_files) if total_files else "[QUANTIDADE]"
    return quantidade, volume, total_files


def _hash_description(data: ReportData) -> str:
    return (
        "Os valores de integridade criptográfica (hashes) referentes aos arquivos contidos nos "
        "pacotes criptografados (formato GPG) foram devidamente ratificados pela respectiva "
        "provedora de aplicação. A referida documentação técnica encontra-se encartada aos "
        f"autos do procedimento {data.procedimento_referenciado or '[PROCEDIMENTO REFERENCIADO]'}, "
        f"sob o ID de documento {data.id_documento or '[ID DO DOCUMENTO]'}."
    )


def render_report(data: ReportData) -> str:
    quantidade, volume, _ = _summary(data)

    if data.hash_rows:
        table_lines = ["| Arquivo | Hash |", "|---|---|"]
        for row in data.hash_rows:
            hash_value = row.sha256 or "(arquivo nao encontrado no destino)"
            table_lines.append(f"| {row.file_name} | {hash_value} |")
        table = "\n".join(table_lines)
    else:
        table = "| [nome-do-arquivo] | [hash] |"

    return f"""# TERMO DE RECEBIMENTO E IDENTIFICAÇÃO DE EVIDÊNCIA TELEMÁTICA

## 1. IDENTIFICAÇÃO DO CASO

- **Número do Processo Judicial:** {data.numero_processo or "[NÚMERO DO PROCESSO JUDICIAL]"}
- **PIC/Inquérito:** {data.numero_pic or "[NÚMERO DO PIC/INQUÉRITO]"}
- **Órgão de Execução:** {data.orgao_execucao or "[ÓRGÃO DE EXECUÇÃO]"}
- **Provedor de Aplicação:** Apple Inc.

## 2. DADOS DA RECEPÇÃO

- **Data/Hora:** {data.data or "[DATA]"}, às {data.hora or "[HORA]"}
- **Responsável:** {data.responsavel or "[NOME DO RESPONSÁVEL]"}
- **Matrícula:** {data.matricula or "[NÚMERO DE MATRÍCULA]"}
- **Forma de Recebimento:**
  - {_checkbox(data.recebido_portal)} Portal do Provedor
  - {_checkbox(data.recebido_email)} E-mail Oficial

## 3. ESPECIFICAÇÕES TÉCNICAS (DADOS BRUTOS)

- **Volume Total (Bytes/GB/TB):** {volume}
- **Quantidade de Arquivos:** {quantidade}
- **Código(s) Hash (Original/Nativo):** {_hash_description(data)}

**Códigos Hash (Gerados pelo órgão):**

{table}

- **Algoritmo utilizado:** SHA-256

## 4. REGISTRO DE CUSTÓDIA E INTEGRIDADE

Os dados acima foram recebidos e processados conforme os requisitos de Integridade (permanência inalterada) e Autenticidade (vínculo ao fato investigado). O material foi imediatamente replicado para backup criptografado institucional.

## 5. DISPONIBILIZAÇÃO

- {_checkbox(data.disponibiliza_original)} Cópia da Aquisição Forense Original (Dados Brutos + Hashes + Metadados)
- {_checkbox(data.disponibiliza_processada)} Cópia Processada/Indexada
"""


# ---------------------------------------------------------------- HTML (clipboard)

def _esc(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def render_report_html(data: ReportData) -> str:
    """Fragmento HTML equivalente ao markdown de render_report(), com fonte Arial 12
    e espaçamento de linha 1.5 embutidos via estilo inline, para colar formatado no Word."""
    quantidade, volume, _ = _summary(data)

    if data.hash_rows:
        rows_html = "".join(
            "<tr>"
            f'<td style="border:1px solid #000;padding:4px 8px;">{_esc(row.file_name)}</td>'
            f'<td style="border:1px solid #000;padding:4px 8px;">'
            f'{_esc(row.sha256) if row.sha256 else "(arquivo nao encontrado no destino)"}</td>'
            "</tr>"
            for row in data.hash_rows
        )
    else:
        rows_html = (
            '<tr><td style="border:1px solid #000;padding:4px 8px;">[nome-do-arquivo]</td>'
            '<td style="border:1px solid #000;padding:4px 8px;">[hash]</td></tr>'
        )

    body_style = f"font-family:{FONT_NAME}, sans-serif; font-size:{FONT_SIZE_PT}pt; line-height:{LINE_SPACING};"
    heading_style = f"font-family:{FONT_NAME}, sans-serif;"

    return f"""<div style="{body_style}">
<h1 style="{heading_style}">TERMO DE RECEBIMENTO E IDENTIFICAÇÃO DE EVIDÊNCIA TELEMÁTICA</h1>

<h2 style="{heading_style}">1. IDENTIFICAÇÃO DO CASO</h2>
<ul>
<li><strong>Número do Processo Judicial:</strong> {_esc(data.numero_processo or "[NÚMERO DO PROCESSO JUDICIAL]")}</li>
<li><strong>PIC/Inquérito:</strong> {_esc(data.numero_pic or "[NÚMERO DO PIC/INQUÉRITO]")}</li>
<li><strong>Órgão de Execução:</strong> {_esc(data.orgao_execucao or "[ÓRGÃO DE EXECUÇÃO]")}</li>
<li><strong>Provedor de Aplicação:</strong> Apple Inc.</li>
</ul>

<h2 style="{heading_style}">2. DADOS DA RECEPÇÃO</h2>
<ul>
<li><strong>Data/Hora:</strong> {_esc(data.data or "[DATA]")}, às {_esc(data.hora or "[HORA]")}</li>
<li><strong>Responsável:</strong> {_esc(data.responsavel or "[NOME DO RESPONSÁVEL]")}</li>
<li><strong>Matrícula:</strong> {_esc(data.matricula or "[NÚMERO DE MATRÍCULA]")}</li>
</ul>
<p><strong>Forma de Recebimento:</strong></p>
<ul>
<li>{_checkbox_html(data.recebido_portal)} Portal do Provedor</li>
<li>{_checkbox_html(data.recebido_email)} E-mail Oficial</li>
</ul>

<h2 style="{heading_style}">3. ESPECIFICAÇÕES TÉCNICAS (DADOS BRUTOS)</h2>
<ul>
<li><strong>Volume Total (Bytes/GB/TB):</strong> {volume}</li>
<li><strong>Quantidade de Arquivos:</strong> {quantidade}</li>
<li><strong>Código(s) Hash (Original/Nativo):</strong> {_esc(_hash_description(data))}</li>
</ul>

<p><strong>Códigos Hash (Gerados pelo órgão):</strong></p>
<table style="border-collapse:collapse;">
<tr>
<th style="border:1px solid #000;padding:4px 8px;text-align:left;">Arquivo</th>
<th style="border:1px solid #000;padding:4px 8px;text-align:left;">Hash</th>
</tr>
{rows_html}
</table>
<p><strong>Algoritmo utilizado:</strong> SHA-256</p>

<h2 style="{heading_style}">4. REGISTRO DE CUSTÓDIA E INTEGRIDADE</h2>
<p>Os dados acima foram recebidos e processados conforme os requisitos de Integridade (permanência inalterada) e Autenticidade (vínculo ao fato investigado). O material foi imediatamente replicado para backup criptografado institucional.</p>

<h2 style="{heading_style}">5. DISPONIBILIZAÇÃO</h2>
<ul>
<li>{_checkbox_html(data.disponibiliza_original)} Cópia da Aquisição Forense Original (Dados Brutos + Hashes + Metadados)</li>
<li>{_checkbox_html(data.disponibiliza_processada)} Cópia Processada/Indexada</li>
</ul>
</div>"""


# ---------------------------------------------------------------------- .docx

def build_docx(data: ReportData, path: Path) -> None:
    """Gera um .docx real (Arial 12, espaçamento 1.5, tabela nativa) com o termo."""
    from docx import Document
    from docx.shared import Pt

    document = Document()

    normal = document.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(FONT_SIZE_PT)
    normal.paragraph_format.line_spacing = LINE_SPACING

    def run_style(run, bold=False):
        run.bold = bold
        run.font.name = FONT_NAME
        run.font.size = Pt(FONT_SIZE_PT)

    def add_heading(text, size):
        p = document.add_paragraph()
        p.paragraph_format.line_spacing = LINE_SPACING
        p.paragraph_format.space_after = Pt(6)
        run_style(p.add_run(text), bold=True)
        p.runs[0].font.size = Pt(size)
        return p

    def add_bullet(label, value):
        p = document.add_paragraph(style="List Bullet")
        p.paragraph_format.line_spacing = LINE_SPACING
        run_style(p.add_run(f"{label}: "), bold=True)
        run_style(p.add_run(value))
        return p

    def add_subbullet(checked, text):
        p = document.add_paragraph(style="List Bullet 2")
        p.paragraph_format.line_spacing = LINE_SPACING
        marker = "☑ " if checked else "☐ "
        run_style(p.add_run(marker + text))
        return p

    def add_paragraph(text="", bold_prefix=None):
        p = document.add_paragraph()
        p.paragraph_format.line_spacing = LINE_SPACING
        if bold_prefix:
            run_style(p.add_run(bold_prefix), bold=True)
        if text:
            run_style(p.add_run(text))
        return p

    quantidade, volume, _ = _summary(data)

    add_heading("TERMO DE RECEBIMENTO E IDENTIFICAÇÃO DE EVIDÊNCIA TELEMÁTICA", 15)

    add_heading("1. IDENTIFICAÇÃO DO CASO", 13)
    add_bullet("Número do Processo Judicial", data.numero_processo or "[NÚMERO DO PROCESSO JUDICIAL]")
    add_bullet("PIC/Inquérito", data.numero_pic or "[NÚMERO DO PIC/INQUÉRITO]")
    add_bullet("Órgão de Execução", data.orgao_execucao or "[ÓRGÃO DE EXECUÇÃO]")
    add_bullet("Provedor de Aplicação", "Apple Inc.")

    add_heading("2. DADOS DA RECEPÇÃO", 13)
    add_bullet("Data/Hora", f"{data.data or '[DATA]'}, às {data.hora or '[HORA]'}")
    add_bullet("Responsável", data.responsavel or "[NOME DO RESPONSÁVEL]")
    add_bullet("Matrícula", data.matricula or "[NÚMERO DE MATRÍCULA]")
    add_paragraph(bold_prefix="Forma de Recebimento:")
    add_subbullet(data.recebido_portal, "Portal do Provedor")
    add_subbullet(data.recebido_email, "E-mail Oficial")

    add_heading("3. ESPECIFICAÇÕES TÉCNICAS (DADOS BRUTOS)", 13)
    add_bullet("Volume Total (Bytes/GB/TB)", volume)
    add_bullet("Quantidade de Arquivos", quantidade)
    add_bullet("Código(s) Hash (Original/Nativo)", _hash_description(data))

    add_paragraph(bold_prefix="Códigos Hash (Gerados pelo órgão):")
    rows = data.hash_rows or [HashRow(file_name="[nome-do-arquivo]", sha256="[hash]")]
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("Arquivo", "Hash")):
        run_style(cell.paragraphs[0].add_run(text), bold=True)
    for row in rows:
        hash_display = row.sha256 or "(arquivo nao encontrado no destino)"
        cells = table.add_row().cells
        for cell, text in zip(cells, (row.file_name, hash_display)):
            run_style(cell.paragraphs[0].add_run(text))
    add_bullet("Algoritmo utilizado", "SHA-256")

    add_heading("4. REGISTRO DE CUSTÓDIA E INTEGRIDADE", 13)
    add_paragraph(
        "Os dados acima foram recebidos e processados conforme os requisitos de Integridade "
        "(permanência inalterada) e Autenticidade (vínculo ao fato investigado). O material foi "
        "imediatamente replicado para backup criptografado institucional."
    )

    add_heading("5. DISPONIBILIZAÇÃO", 13)
    add_subbullet(data.disponibiliza_original, "Cópia da Aquisição Forense Original (Dados Brutos + Hashes + Metadados)")
    add_subbullet(data.disponibiliza_processada, "Cópia Processada/Indexada")

    document.save(str(path))
